"""La hoja «Gantt» del export: celdas pintadas que coinciden con el cronograma."""

from io import BytesIO

from openpyxl import load_workbook
from sqlmodel import Session

from app.schemas import TareaIn
from app.services import arbol as arbol_service
from app.services import exportar_excel
from app.services import vista as vista_service

from .test_exportar_excel import poblar

_FIJAS = 4  # Código, Tarea, Inicio, Fin: las columnas de día arrancan después


def hoja_gantt(session, proyecto):
    contenido = exportar_excel.a_excel(session, proyecto.id)
    return load_workbook(BytesIO(contenido))["Gantt"]


def fila_de(hoja, titulo):
    for fila in hoja.iter_rows(min_row=3):
        if fila[1].value == titulo:
            return fila
    raise AssertionError(f"no está la fila «{titulo}»")


def columnas_pintadas(fila):
    """Índices (0-based sobre la fila) de las celdas con relleno de barra."""
    return [
        i for i, celda in enumerate(fila)
        if i >= _FIJAS and celda.fill.fgColor.rgb not in (None, "00000000")
    ]


def test_la_barra_coincide_con_el_cronograma(session: Session, proyecto):
    poblar(session, proyecto)
    datos = vista_service.armar(session, proyecto.id)
    hoja = hoja_gantt(session, proyecto)

    relevamiento = next(f for f in datos.todas_las_filas if f.tarea.titulo == "Relevamiento")
    pintadas = columnas_pintadas(fila_de(hoja, "Relevamiento"))
    # 5 días hábiles = 5 celdas pintadas, y la primera es el día de inicio.
    assert len(pintadas) == 5
    primer_dia = hoja.cell(row=2, column=pintadas[0] + 1).value
    assert str(relevamiento.inicio.day) == str(primer_dia)


def test_la_critica_se_pinta_de_rojo_y_la_comun_de_azul(session: Session, proyecto):
    poblar(session, proyecto)
    hoja = hoja_gantt(session, proyecto)

    comun = fila_de(hoja, "Relevamiento")
    critica = fila_de(hoja, "Análisis")
    assert comun[columnas_pintadas(comun)[0]].fill.fgColor.rgb == "004C8DFF"
    assert critica[columnas_pintadas(critica)[0]].fill.fgColor.rgb == "00FF6B5E"


def test_la_etapa_lleva_gris_negrita_y_envuelve_a_sus_hijas(session: Session, proyecto):
    poblar(session, proyecto)
    hoja = hoja_gantt(session, proyecto)

    etapa = fila_de(hoja, "Etapa 1")
    pintadas = columnas_pintadas(etapa)
    assert etapa[pintadas[0]].fill.fgColor.rgb == "0096A0B5"
    assert etapa[0].font.bold and etapa[1].font.bold

    # La envolvente cubre al menos lo que cubre cada hija.
    for hija in ("Relevamiento", "Análisis"):
        assert set(columnas_pintadas(fila_de(hoja, hija))) <= set(pintadas)


def test_el_hito_es_un_rombo_sin_relleno(session: Session, proyecto):
    poblar(session, proyecto)
    hoja = hoja_gantt(session, proyecto)

    hito = fila_de(hoja, "Hito: etapa cerrada")
    assert columnas_pintadas(hito) == []
    rombos = [c.value for c in hito[_FIJAS:] if c.value == "◆"]
    assert rombos == ["◆"]


def test_un_proyecto_largo_sale_por_semana(session: Session, proyecto):
    arbol_service.agregar_al_final(
        session, proyecto.id, TareaIn(titulo="Acompañamiento anual", duracion=260)
    )
    hoja = hoja_gantt(session, proyecto)

    etiquetas = [c.value for c in hoja[2][_FIJAS:] if c.value is not None]
    assert all(str(e).startswith("S") for e in etiquetas)  # S01, S02… no días
    assert len(etiquetas) < 80  # ~52 semanas, no ~260 columnas de día
    assert len(columnas_pintadas(fila_de(hoja, "Acompañamiento anual"))) == len(etiquetas)


def test_un_proyecto_vacio_no_rompe_la_hoja(session: Session, proyecto):
    hoja = hoja_gantt(session, proyecto)
    assert hoja.cell(row=2, column=1).value == "Código"
