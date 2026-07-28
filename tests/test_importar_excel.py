"""Import de planillas, probado contra la planilla real que originó la función."""

from datetime import date
from pathlib import Path

from sqlmodel import Session

from app.services import importar as importar_service
from app.services import importar_aplicar
from app.services import importar_excel
from app.services import schedule as schedule_service
from app.services import tasks as tasks_service

PLANILLA = Path(__file__).parent / "fixtures" / "gantt-traspaso.xlsx"


def contenido() -> bytes:
    return PLANILLA.read_bytes()


def test_lista_las_hojas():
    assert importar_excel.hojas(contenido()) == ["Sheet1", "Sheet2"]


def test_por_default_lee_la_ultima_hoja_con_datos():
    assert "Sheet2" in importar_excel.leer(contenido()).origen


def test_saltea_las_hojas_sin_columnas_de_tareas():
    """Una hoja de notas al final no puede ganarle a la del plan."""
    from io import BytesIO

    from openpyxl import Workbook

    libro = Workbook()
    libro.active.title = "Plan"
    for fila in [["WBS", "Tarea", "Días"], ["1", "Arrancar", 2]]:
        libro.active.append(fila)
    notas = libro.create_sheet("Notas")
    notas.append(["Acordarse de pedir los accesos"])
    buffer = BytesIO()
    libro.save(buffer)

    imp = importar_excel.leer(buffer.getvalue())
    assert "Plan" in imp.origen
    assert [f.titulo for f in imp.filas] == ["Arrancar"]


def test_lee_la_estructura_completa():
    imp = importar_excel.leer(contenido(), "Sheet2")
    tipos = [f.tipo for f in imp.filas]
    assert tipos.count("fase") == 5
    assert tipos.count("hito") == 5
    assert len(imp.filas) == 41


def test_la_jerarquia_sale_del_wbs():
    imp = importar_excel.leer(contenido(), "Sheet2")
    por_wbs = {f.wbs: f for f in imp.filas if f.wbs}
    assert por_wbs["1"].nivel == 0 and por_wbs["1"].tipo == "fase"
    assert por_wbs["1.1"].nivel == 1
    assert por_wbs["1.1"].duracion == 5


def test_los_hitos_quedan_con_duracion_cero():
    imp = importar_excel.leer(contenido(), "Sheet2")
    hitos = [f for f in imp.filas if f.tipo == "hito"]
    assert all(f.duracion == 0 for f in hitos)
    assert hitos[0].titulo.startswith("Hito: Equipo conformado")


def test_repara_los_wbs_que_excel_convirtio_en_fecha():
    imp = importar_excel.leer(contenido(), "Sheet2")
    codigos = {f.wbs for f in imp.filas}
    assert {"4.6", "4.7", "5.7", "5.8", "5.9", "5.10"} <= codigos
    assert not any("2026" in (f.wbs or "") for f in imp.filas)
    assert any("Excel había convertido" in a for a in imp.avisos)


def test_repara_las_predecesoras_convertidas_en_fecha():
    imp = importar_excel.leer(contenido(), "Sheet2")
    por_wbs = {f.wbs: f for f in imp.filas if f.wbs}
    assert por_wbs["5.1"].predecesoras == ["4.7"]
    assert por_wbs["5.7"].predecesoras == ["5.1"]


def test_importa_la_criticidad_marcada_en_la_planilla():
    """La columna «Crít.» del usuario entra como criticidad de negocio."""
    imp = importar_excel.leer(contenido(), "Sheet2")
    criticas = [f.wbs for f in imp.filas if f.critica]
    assert criticas, "la planilla marca varias tareas con Sí"
    assert "1.2" in criticas
    assert not any(f.critica for f in imp.filas if f.wbs == "1.1")


def test_avisa_de_las_filas_sin_wbs():
    imp = importar_excel.leer(contenido(), "Sheet2")
    assert any("no tiene WBS" in a for a in imp.avisos)


def test_descarta_la_leyenda_del_final():
    imp = importar_excel.leer(contenido(), "Sheet2")
    assert not any("Editá las celdas" in f.titulo for f in imp.filas)


def planilla_en_memoria(filas: list[list]) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook

    libro = Workbook()
    for fila in filas:
        libro.active.append(fila)
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def test_una_planilla_sin_cabecera_avisa_y_no_rompe():
    imp = importar_excel.leer(planilla_en_memoria([["cualquier", "cosa"], [1, 2]]))
    assert imp.filas == []
    assert any("encabezados" in a for a in imp.avisos)


def test_el_lag_escrito_en_la_planilla_se_respeta(session: Session):
    """`1+2` en la columna Predec. significa lo mismo que en la grilla."""
    contenido = planilla_en_memoria(
        [
            ["WBS", "Tarea", "Predec.", "Días"],
            ["1", "A", "", 3],
            ["2", "B", "1+2", 1],
            ["3", "C", "1-1", 1],
        ]
    )
    imp = importar_excel.leer(contenido)
    proyecto, avisos = importar_aplicar.aplicar(session, "Lag", date(2026, 1, 5), imp)
    assert avisos == []

    plan = schedule_service.calcular(session, proyecto.id)
    por_titulo = {t.titulo: t.id for t, _ in tasks_service.arbol(session, proyecto.id)}
    assert plan.get(por_titulo["A"]).fin == date(2026, 1, 7)
    assert plan.get(por_titulo["B"]).inicio == date(2026, 1, 12)  # +2 días hábiles
    assert plan.get(por_titulo["C"]).inicio == date(2026, 1, 7)   # solapa 1


def test_una_planilla_minima_alcanza_con_wbs_y_tarea():
    imp = importar_excel.leer(
        planilla_en_memoria([["WBS", "Tarea", "Días"], ["1", "Arrancar", 3]])
    )
    assert len(imp.filas) == 1
    assert imp.filas[0].titulo == "Arrancar" and imp.filas[0].duracion == 3


def test_aplicar_crea_el_proyecto_con_arbol_y_dependencias(session: Session):
    imp = importar_excel.leer(contenido(), "Sheet2")
    proyecto, avisos = importar_aplicar.aplicar(
        session, "Traspaso operativo", date(2026, 8, 3), imp
    )

    arbol = tasks_service.arbol(session, proyecto.id)
    assert len(arbol) == 41
    titulos = {t.titulo: (t, n) for t, n in arbol}
    assert titulos["Etapa 1 — Conformación del Equipo"][1] == 0
    assert titulos["Definir denominación y estructura del nuevo sector"][1] == 1

    # la única dependencia que no se puede crear es la que apunta al WBS inexistente
    assert sum("3.2" in a for a in avisos) == 1


def test_el_cronograma_del_import_calcula_sin_errores(session: Session):
    imp = importar_excel.leer(contenido(), "Sheet2")
    proyecto, _ = importar_aplicar.aplicar(
        session, "Traspaso", date(2026, 8, 3), imp
    )
    plan, error = schedule_service.calcular_seguro(session, proyecto.id)
    assert error is None
    assert plan.inicio == date(2026, 8, 3)
    assert plan.fin > plan.inicio
