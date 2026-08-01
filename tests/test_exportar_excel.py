"""Export a Excel. El criterio de esta fase es la **ida y vuelta**.

Exportar un proyecto y reimportar ese mismo archivo tiene que reproducirlo. Es el
test más fuerte que se puede pedir acá: si una columna se escribe con un nombre que
el importador no entiende, o se pierde un dato al escribirlo, esto se cae.
"""

from io import BytesIO

from openpyxl import load_workbook
from sqlmodel import Session

from app.models import Ambito
from app.schemas import ProyectoIn, TareaIn
from app.services import arbol as arbol_service
from app.services import exportar_excel
from app.services import importar_aplicar
from app.services import importar_excel
from app.services import predecesoras as predecesoras_service
from app.services import projects as projects_service
from app.services import riesgos as riesgos_service
from app.services import tasks as tasks_service
from app.services import vista as vista_service


def poblar(session, proyecto):
    """Un proyecto con de todo: jerarquía, hito, dependencias, ámbitos y pesos."""
    etapa = arbol_service.agregar_al_final(
        session, proyecto.id, TareaIn(titulo="Etapa 1", peso=70)
    )
    sub_a = arbol_service.agregar_al_final(
        session, proyecto.id,
        TareaIn(titulo="Relevamiento", duracion=5, responsable="Ariel", peso=60,
                duracion_optimista=3, duracion_pesimista=9),
    )
    sub_b = arbol_service.agregar_al_final(
        session, proyecto.id,
        TareaIn(titulo="Análisis", duracion=3, critica=True, peso=40),
    )
    for hija in (sub_a, sub_b):
        arbol_service.indentar(session, hija.id)
        tasks_service.mover(session, hija.id, etapa.id)

    hito = arbol_service.agregar_al_final(
        session, proyecto.id, TareaIn(titulo="Hito: etapa cerrada", duracion=0, peso=10)
    )
    acomp = arbol_service.agregar_al_final(
        session, proyecto.id,
        TareaIn(titulo="Acompañamiento", duracion=20, ambito=Ambito.seguimiento, peso=20),
    )
    predecesoras_service.guardar(session, proyecto.id, sub_b.id, f"{sub_a.codigo}+2")
    predecesoras_service.guardar(session, proyecto.id, hito.id, sub_b.codigo)
    return {"etapa": etapa, "sub_a": sub_a, "sub_b": sub_b, "hito": hito, "acomp": acomp}


def test_el_libro_se_genera_y_abre(session: Session, proyecto):
    poblar(session, proyecto)
    contenido = exportar_excel.a_excel(session, proyecto.id)

    libro = load_workbook(BytesIO(contenido))
    assert libro.sheetnames == ["Plan", "Gantt"]
    cabeceras = [c.value for c in libro["Plan"][1]]
    assert cabeceras[:5] == ["WBS", "Tarea", "Resp.", "Predec.", "Días"]


def test_un_proyecto_que_no_existe_no_rompe(session: Session):
    assert exportar_excel.a_excel(session, 9999) is None


def test_el_resumen_no_lleva_duracion_propia(session: Session, proyecto):
    poblar(session, proyecto)
    hoja = load_workbook(BytesIO(exportar_excel.a_excel(session, proyecto.id)))["Plan"]
    fila = next(f for f in hoja.iter_rows(min_row=2, values_only=True) if f[1] == "Etapa 1")
    assert fila[4] is None  # Días: lo calcula el rollup


def test_las_predecesoras_salen_en_notacion_wbs(session: Session, proyecto):
    creadas = poblar(session, proyecto)
    hoja = load_workbook(BytesIO(exportar_excel.a_excel(session, proyecto.id)))["Plan"]
    fila = next(f for f in hoja.iter_rows(min_row=2, values_only=True) if f[1] == "Análisis")
    assert fila[3] == f"{creadas['sub_a'].codigo}+2"


def test_el_riesgo_se_marca(session: Session, proyecto):
    creadas = poblar(session, proyecto)
    riesgos_service.marcar_tarea(session, proyecto.id, creadas["sub_a"].id, True)

    hoja = load_workbook(BytesIO(exportar_excel.a_excel(session, proyecto.id)))["Plan"]
    filas = {f[1]: f for f in hoja.iter_rows(min_row=2, values_only=True)}
    assert filas["Relevamiento"][12] == "Sí"
    assert not filas["Análisis"][12]  # openpyxl guarda la cadena vacía como celda vacía


# --- el contrato de la fase: ida y vuelta ---

def test_exportar_y_reimportar_reproduce_el_proyecto(session: Session, proyecto):
    poblar(session, proyecto)
    original = vista_service.armar(session, proyecto.id)

    contenido = exportar_excel.a_excel(session, proyecto.id)
    importacion = importar_excel.leer(contenido)
    copia, _ = importar_aplicar.aplicar(
        session, "Copia", proyecto.fecha_inicio, importacion
    )
    vuelta = vista_service.armar(session, copia.id)

    def retrato(datos):
        return [
            (
                f.tarea.codigo, f.tarea.titulo, f.nivel, f.tarea.duracion,
                f.responsable.nombre if f.responsable else "",
                f.tarea.critica, f.tarea.ambito.value,
                f.tarea.peso, f.tarea.duracion_optimista, f.tarea.duracion_pesimista,
                f.predecesoras_texto, f.es_hito, f.es_resumen,
            )
            for f in datos.todas_las_filas
        ]

    assert retrato(vuelta) == retrato(original)


def test_la_vuelta_calcula_las_mismas_fechas(session: Session, proyecto):
    """Si las dependencias sobrevivieron bien, el cronograma tiene que dar igual."""
    poblar(session, proyecto)
    original = vista_service.armar(session, proyecto.id)

    importacion = importar_excel.leer(exportar_excel.a_excel(session, proyecto.id))
    copia, _ = importar_aplicar.aplicar(
        session, "Copia", proyecto.fecha_inicio, importacion
    )
    vuelta = vista_service.armar(session, copia.id)

    assert vuelta.resumen.inicio == original.resumen.inicio
    assert vuelta.resumen.fin == original.resumen.fin
    assert vuelta.resumen.dias_habiles == original.resumen.dias_habiles
    assert vuelta.avance_ponderado == original.avance_ponderado


def test_la_vuelta_no_arrastra_avisos_de_reparacion(session: Session, proyecto):
    """Un archivo que generamos nosotros no debería necesitar arreglos al leerlo."""
    poblar(session, proyecto)
    importacion = importar_excel.leer(exportar_excel.a_excel(session, proyecto.id))
    assert importacion.avisos == []


def test_las_columnas_calculadas_no_ensucian_la_vuelta(session: Session, proyecto):
    """Fechas, holgura y avance se exportan para leer; al reimportar los pone el motor."""
    poblar(session, proyecto)
    importacion = importar_excel.leer(exportar_excel.a_excel(session, proyecto.id))
    # El lector las reconoce como fechas declaradas, pero no fija nada con ellas.
    assert all(f.inicio_declarado is not None for f in importacion.filas if f.wbs)

    copia, _ = importar_aplicar.aplicar(
        session, "Copia", proyecto.fecha_inicio + __import__("datetime").timedelta(days=7),
        importacion,
    )
    vuelta = vista_service.armar(session, copia.id)
    # Arrancando una semana después, el cronograma se corre: las fechas de la
    # planilla no lo anclaron.
    assert vuelta.resumen.inicio > vista_service.armar(session, proyecto.id).resumen.inicio


def test_un_proyecto_vacio_tambien_va_y_vuelve(session: Session):
    from datetime import date

    vacio = projects_service.crear(
        session, ProyectoIn(nombre="Vacío", fecha_inicio=date(2026, 1, 5))
    )
    importacion = importar_excel.leer(exportar_excel.a_excel(session, vacio.id))
    assert importacion.filas == []
