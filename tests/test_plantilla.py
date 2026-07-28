"""La planilla modelo define el formato de importación, así que tiene que poder
volver a entrar por el importador sin perder nada."""

from io import BytesIO

from openpyxl import load_workbook

from app.services import importar_excel
from app.services import plantilla as plantilla_service


def libro():
    return load_workbook(BytesIO(plantilla_service.construir()))


def test_tiene_las_dos_hojas():
    assert libro().sheetnames == ["Plan", "Cómo se completa"]


def test_los_encabezados_son_los_que_el_importador_espera():
    hoja = libro()["Plan"]
    encabezados = [c.value for c in hoja[2] if c.value]
    assert encabezados == ["WBS", "Tarea", "Resp.", "Predec.", "Días", "_tipo"]


def test_las_columnas_de_codigo_son_texto():
    """Sin esto Excel convierte 4.6 en una fecha y rompe la jerarquía."""
    hoja = libro()["Plan"]
    assert hoja["A3"].number_format == "@"   # WBS
    assert hoja["D3"].number_format == "@"   # Predec.
    # también las filas vacías que el usuario va a completar
    assert hoja["A50"].number_format == "@"


def test_la_plantilla_se_puede_importar_de_vuelta():
    imp = importar_excel.leer(plantilla_service.construir(), "Plan")
    assert imp.avisos == []
    por_wbs = {f.wbs: f for f in imp.filas}
    assert set(por_wbs) == {"1", "1.1", "1.2", "1.3", "2", "2.1", "2.2"}
    assert por_wbs["1"].tipo == "fase"
    assert por_wbs["1.3"].tipo == "hito" and por_wbs["1.3"].duracion == 0
    assert por_wbs["1.2"].predecesoras == ["1.1"]
    assert por_wbs["2.2"].predecesoras == ["2.1-2"]
    assert por_wbs["1.1"].responsable == "Ariel"


def test_el_ejemplo_de_la_plantilla_calcula_un_cronograma(session, proyecto):
    from datetime import date

    from app.services import importar_aplicar
    from app.services import schedule as schedule_service

    imp = importar_excel.leer(plantilla_service.construir(), "Plan")
    nuevo, avisos = importar_aplicar.aplicar(session, "Modelo", date(2026, 1, 5), imp)
    plan, error = schedule_service.calcular_seguro(session, nuevo.id)
    assert error is None
    assert avisos == []
    assert plan.inicio == date(2026, 1, 5)
