"""Export del proyecto a `.xlsx`, pensado para poder volver a entrar por el importador.

**El contrato es la ida y vuelta:** exportar un proyecto y reimportar ese mismo archivo
tiene que reproducirlo. Por eso las cabeceras son exactamente las que
`importar_excel` sabe leer, y las columnas calculadas van al final y marcadas: si
alguien reimporta el archivo, las fechas y el avance se ignoran —los calcula el motor—
y no ensucian nada.

Se genera `.xlsx`, no `.xls`: el binario viejo necesitaría una librería abandonada y
Excel abre `.xlsx` sin chistar desde 2007.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlmodel import Session

from ..models import ETIQUETA_AMBITO
from . import riesgos as riesgos_service
from . import vista as vista_service
from .gantt_vista import avance

# (encabezado, ancho). El orden es el de la grilla, para que se lea igual.
_COLUMNAS = [
    ("WBS", 10), ("Tarea", 52), ("Resp.", 16), ("Predec.", 16), ("Días", 7),
    ("Opt", 6), ("Pes", 6), ("Ámbito", 13), ("Peso", 7), ("Crít.", 7), ("_tipo", 8),
    ("Estado", 14), ("Riesgo", 8), ("Inicio", 12), ("Fin", 12),
    ("% proyecto", 11), ("Avance", 9), ("Holgura", 9),
]
# De acá en adelante lo calcula el motor: se exporta para leer, no para reimportar.
_PRIMERA_CALCULADA = 13

_CABECERA = PatternFill("solid", fgColor="1A1F2B")
_CALCULADA = PatternFill("solid", fgColor="F0F2F6")


def a_excel(session: Session, project_id: int) -> bytes | None:
    """Un libro con una hoja: la grilla completa del proyecto."""
    from . import projects as projects_service

    proyecto = projects_service.obtener(session, project_id)
    if proyecto is None:
        return None

    datos = vista_service.armar(session, project_id)
    con_riesgo = riesgos_service.ids_con_riesgo(session, project_id)

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Plan"
    _encabezar(hoja)

    for numero, fila in enumerate(datos.todas_las_filas, start=2):
        for columna, valor in enumerate(_valores(fila, con_riesgo), start=1):
            celda = hoja.cell(row=numero, column=columna, value=valor)
            if columna >= _PRIMERA_CALCULADA:
                celda.fill = _CALCULADA
            if fila.es_resumen:
                celda.font = Font(bold=True)
        # La sangría del título es la única pista visual de la jerarquía en Excel.
        hoja.cell(row=numero, column=2).alignment = Alignment(indent=fila.nivel)

    hoja.freeze_panes = "C2"
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _encabezar(hoja) -> None:
    for indice, (titulo, ancho) in enumerate(_COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=indice, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = _CABECERA
        hoja.column_dimensions[get_column_letter(indice)].width = ancho


def _valores(fila, con_riesgo: set[int]) -> list:
    tarea = fila.tarea
    return [
        tarea.codigo,
        tarea.titulo,
        tarea.responsable,
        fila.predecesoras_texto,
        # Un resumen no tiene duración propia: la escribe el rollup de sus hijas.
        None if fila.es_resumen else tarea.duracion,
        tarea.duracion_optimista,
        tarea.duracion_pesimista,
        ETIQUETA_AMBITO[tarea.ambito],
        tarea.peso,
        "Sí" if tarea.critica else "",
        _tipo(fila),
        fila.estado.nombre if fila.estado else "",
        "Sí" if (tarea.id or 0) in con_riesgo else "",
        fila.inicio,
        fila.fin,
        round(fila.peso_absoluto, 1),
        avance(fila),
        None if fila.es_resumen else fila.holgura,
    ]


def _tipo(fila) -> str:
    if fila.es_hito:
        return "hito"
    return "fase" if fila.es_resumen else "tarea"
