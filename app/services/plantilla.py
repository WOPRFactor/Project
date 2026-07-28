"""Genera la planilla modelo que define el formato de importación.

Es la forma más honesta de fijar el estándar: en vez de documentar columnas en un
texto que nadie lee, se descarga un archivo que ya tiene la estructura correcta.

Detalle que importa: WBS y Predec. se marcan como **texto** en el formato de celda.
Sin eso, Excel convierte `4.6` en una fecha y rompe la jerarquía — es exactamente
el problema que aparece en las planillas hechas a mano.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNAS = [
    ("WBS", 10, "Código jerárquico. 1, 1.1, 1.2, 2… El punto marca el nivel."),
    ("Tarea", 52, "Qué hay que hacer."),
    ("Resp.", 16, "Quién. Opcional."),
    ("Predec.", 14, "De qué depende, por WBS. 1.3 · varias con coma · 1.3+2 espera 2 días."),
    ("Días", 8, "Duración en días hábiles. 0 = hito."),
    ("_tipo", 10, "fase, tarea o hito. Si se omite, se deduce."),
]

EJEMPLO = [
    ("1", "Relevamiento", "", "", None, "fase"),
    ("1.1", "Entrevistas con el cliente", "Ariel", "", 4, "tarea"),
    ("1.2", "Inventario de activos", "", "1.1", 3, "tarea"),
    ("1.3", "Hito: relevamiento cerrado", "", "1.2", 0, "hito"),
    ("2", "Diseño", "", "", None, "fase"),
    ("2.1", "Arquitectura objetivo", "Ariel", "1.3", 5, "tarea"),
    ("2.2", "Plan de controles", "", "2.1-2", 3, "tarea"),
]

_TEXTO = "@"
_ENCABEZADO = PatternFill("solid", fgColor="1F3864")
_TITULO = PatternFill("solid", fgColor="D9E2F3")


def construir(nombre_proyecto: str = "Proyecto nuevo") -> bytes:
    libro = Workbook()
    _hoja_plan(libro.active, nombre_proyecto)
    _hoja_ayuda(libro.create_sheet("Cómo se completa"))
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _hoja_plan(hoja, nombre_proyecto: str) -> None:
    hoja.title = "Plan"
    hoja["A1"] = nombre_proyecto
    hoja["A1"].font = Font(bold=True, size=13)
    hoja["A1"].fill = _TITULO

    for indice, (titulo, ancho, ayuda) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=2, column=indice, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = _ENCABEZADO
        celda.alignment = Alignment(horizontal="center")
        celda.comment = None
        letra = get_column_letter(indice)
        hoja.column_dimensions[letra].width = ancho

    for numero, fila in enumerate(EJEMPLO, start=3):
        for indice, valor in enumerate(fila, start=1):
            celda = hoja.cell(row=numero, column=indice, value=valor)
            if COLUMNAS[indice - 1][0] in ("WBS", "Predec."):
                celda.number_format = _TEXTO

    # Todo el resto de la columna queda como texto, para que Excel no convierta
    # los WBS que se agreguen después.
    for columna in ("A", "D"):
        for numero in range(len(EJEMPLO) + 3, 400):
            hoja[f"{columna}{numero}"].number_format = _TEXTO

    hoja.freeze_panes = "A3"


def _hoja_ayuda(hoja) -> None:
    hoja["A1"] = "Cómo se completa esta planilla"
    hoja["A1"].font = Font(bold=True, size=13)
    hoja.column_dimensions["A"].width = 14
    hoja.column_dimensions["B"].width = 92

    filas = [("Columna", "Qué va")] + [(t, ayuda) for t, _, ayuda in COLUMNAS]
    filas += [
        ("", ""),
        ("Fechas", "No se cargan: las calcula la app desde la fecha de inicio del "
                   "proyecto, las duraciones y las dependencias."),
        ("Días hábiles", "Se cuentan de lunes a viernes. Los fines de semana se saltean solos."),
        ("Jerarquía", "Una tarea con subtareas no lleva duración: sus fechas envuelven a las hijas."),
        ("Ojo con Excel", "Las columnas WBS y Predec. ya vienen con formato de texto. "
                          "Si las cambiás a General, Excel convierte 4.6 en una fecha."),
    ]
    for numero, (izquierda, derecha) in enumerate(filas, start=3):
        hoja.cell(row=numero, column=1, value=izquierda).font = Font(bold=True)
        celda = hoja.cell(row=numero, column=2, value=derecha)
        celda.alignment = Alignment(wrap_text=True, vertical="top")
