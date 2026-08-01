"""Hoja «Gantt» del export a Excel: el timeline reconstruido con celdas pintadas.

El Gantt de la app es HTML+CSS y no se puede incrustar en un `.xlsx`; esto lo
reconstruye como lo exporta MS Project: una columna angosta por día hábil (o por
semana, si el proyecto es largo) y las barras como celdas rellenas. La geometría
sale de `engine/timeline.construir_grilla` — la misma que pinta la pantalla, así
Excel y app nunca difieren.

Ojo con la cabecera: el reimportador elige la última hoja cuyos encabezados
reconoce. Por eso la primera columna se llama «Código» y no «WBS» — si esta hoja
pareciera importable, la ida y vuelta del export leería el Gantt en vez del Plan.
"""

from __future__ import annotations

from datetime import date, timedelta

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..engine.timeline import Semana, construir_grilla

_FIJAS = [("Código", 10), ("Tarea", 46), ("Inicio", 11), ("Fin", 11)]
# Pasadas estas columnas de día, se pinta por semana: mismo espíritu que
# `timeline.ancho_columna`, que en pantalla achica la columna antes que scrollear.
_UMBRAL_SEMANAL = 180

_CABECERA = PatternFill("solid", fgColor="1A1F2B")
_BARRA = PatternFill("solid", fgColor="4C8DFF")      # --acento
_CRITICA = PatternFill("solid", fgColor="FF6B5E")    # --critico
_RESUMEN = PatternFill("solid", fgColor="96A0B5")    # --tenue
_HOY = PatternFill("solid", fgColor="F5B544")        # --curso
_BLANCA = Font(bold=True, color="FFFFFF")


def agregar_hoja(libro, filas, hoy: date | None = None) -> None:
    """Suma la hoja «Gantt» al libro. `filas` son las de la vista ya armada."""
    hoja = libro.create_sheet("Gantt")
    _fijas(hoja, filas)

    fechas = [(f.inicio, f.fin) for f in filas if f.inicio and f.fin]
    if not fechas:
        return
    grilla = construir_grilla(min(i for i, _ in fechas), max(f for _, f in fechas))

    # Cada columna pintable cubre un tramo [desde, hasta]; por día ambos coinciden.
    if grilla.columnas > _UMBRAL_SEMANAL:
        tramos = [_tramo(s) for s in grilla.semanas]
        etiquetas = [s.etiqueta for s in grilla.semanas]
    else:
        tramos = [(dia, dia) for dia in grilla.dias]
        etiquetas = [str(dia.day) for dia in grilla.dias]

    _cabecera(hoja, grilla, tramos, etiquetas, hoy)
    for numero, fila in enumerate(filas, start=3):
        _pintar(hoja, numero, fila, tramos)

    hoja.freeze_panes = "E3"


def _tramo(semana: Semana) -> tuple[date, date]:
    return semana.lunes, semana.lunes + timedelta(days=4)


def _fijas(hoja, filas) -> None:
    for indice, (titulo, ancho) in enumerate(_FIJAS, start=1):
        celda = hoja.cell(row=2, column=indice, value=titulo)
        celda.font = _BLANCA
        celda.fill = _CABECERA
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    for numero, fila in enumerate(filas, start=3):
        hoja.cell(row=numero, column=1, value=fila.tarea.codigo)
        titulo = hoja.cell(row=numero, column=2, value=fila.tarea.titulo)
        titulo.alignment = Alignment(indent=fila.nivel)
        hoja.cell(row=numero, column=3, value=fila.inicio)
        hoja.cell(row=numero, column=4, value=fila.fin)
        if fila.es_resumen:
            for columna in range(1, 5):
                hoja.cell(row=numero, column=columna).font = Font(bold=True)


def _cabecera(hoja, grilla, tramos, etiquetas, hoy: date | None) -> None:
    base = len(_FIJAS)
    ancho = 1.2 if len(tramos) == grilla.columnas else 3.6

    # Fila 1: los meses, combinados sobre las columnas que cubren.
    desde = base + 1
    for mes, cantidad in _meses(tramos):
        hasta = desde + cantidad - 1
        if hasta > desde:
            hoja.merge_cells(start_row=1, start_column=desde, end_row=1, end_column=hasta)
        celda = hoja.cell(row=1, column=desde, value=mes)
        celda.font = _BLANCA
        celda.fill = _CABECERA
        desde = hasta + 1

    # Fila 2: el día del mes o la semana ISO, uno por columna pintable.
    for indice, etiqueta in enumerate(etiquetas):
        celda = hoja.cell(row=2, column=base + 1 + indice, value=etiqueta)
        celda.font = Font(size=8, color="FFFFFF")
        celda.fill = _CABECERA
        celda.alignment = Alignment(horizontal="center")
        hoja.column_dimensions[get_column_letter(base + 1 + indice)].width = ancho
        if hoy is not None and tramos[indice][0] <= hoy <= tramos[indice][1]:
            celda.fill = _HOY
            celda.font = Font(size=8, bold=True, color="1A1F2B")


def _meses(tramos) -> list[tuple[str, int]]:
    """(etiqueta, cuántas columnas cubre), agrupando por el mes del arranque del tramo."""
    _MESES = [
        "ene", "feb", "mar", "abr", "may", "jun",
        "jul", "ago", "sep", "oct", "nov", "dic",
    ]
    salida: list[tuple[str, int]] = []
    for desde, _ in tramos:
        etiqueta = f"{_MESES[desde.month - 1]} {desde.year}"
        if salida and salida[-1][0] == etiqueta:
            salida[-1] = (etiqueta, salida[-1][1] + 1)
        else:
            salida.append((etiqueta, 1))
    return salida


def _pintar(hoja, numero: int, fila, tramos) -> None:
    if fila.inicio is None or fila.fin is None:
        return
    base = len(_FIJAS)
    for indice, (desde, hasta) in enumerate(tramos):
        if fila.fin < desde or fila.inicio > hasta:
            continue
        celda = hoja.cell(row=numero, column=base + 1 + indice)
        if fila.es_hito:
            # Un hito marca un momento, no un tramo: rombo, sin relleno.
            celda.value = "◆"
            celda.alignment = Alignment(horizontal="center")
        elif fila.es_resumen:
            celda.fill = _RESUMEN
        elif fila.tarea.critica:
            celda.fill = _CRITICA
        else:
            celda.fill = _BARRA
