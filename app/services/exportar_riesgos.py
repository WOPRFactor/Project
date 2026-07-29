"""La hoja «Riesgos» del libro exportado.

Va aparte de la grilla —hoja propia, no columnas más— porque un riesgo no es una
tarea: no tiene fechas ni duración, y meterlo en la misma tabla obligaría a dejar
media planilla vacía. El importador elige la primera hoja con las columnas de
tareas, así que esta hoja no le molesta.

Es de **lectura**: el registro se carga en la app, no se reimporta desde Excel. La
ida y vuelta que el proyecto garantiza es la del cronograma.
"""

from __future__ import annotations

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlmodel import Session

from ..models import ETIQUETA_ESTADO_RIESGO, ETIQUETA_RESPUESTA
from . import riesgos as riesgos_service

_COLUMNAS = [
    ("Riesgo", 46), ("Amenaza a", 26), ("P", 5), ("I", 5), ("Sev.", 6),
    ("Respuesta", 13), ("Plan", 40), ("P res.", 7), ("I res.", 7), ("Sev. res.", 9),
    ("Disparador", 34), ("Responsable", 18), ("Revisar", 12), ("Tarea del plan", 26),
    ("Estado", 14),
]

_CABECERA = PatternFill("solid", fgColor="1A1F2B")


def agregar_hoja(libro, session: Session, project_id: int) -> None:
    """Suma la hoja al libro. Sin riesgos cargados no se agrega nada."""
    panel = riesgos_service.panel(session, project_id)
    if not panel.riesgos:
        return

    hoja = libro.create_sheet("Riesgos")
    for indice, (titulo, ancho) in enumerate(_COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=indice, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = _CABECERA
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    for numero, riesgo in enumerate(panel.riesgos, start=2):
        for columna, valor in enumerate(_valores(riesgo, panel), start=1):
            hoja.cell(row=numero, column=columna, value=valor)

    hoja.freeze_panes = "A2"


def _valores(riesgo, panel) -> list:
    return [
        riesgo.descripcion,
        panel.titulos.get(riesgo.task_id or 0, ""),
        riesgo.probabilidad,
        riesgo.impacto,
        riesgo.severidad,
        ETIQUETA_RESPUESTA[riesgo.respuesta],
        riesgo.mitigacion,
        # Vacío y no repetido: el residual sin estimar no es un dato, es un pendiente.
        riesgo.probabilidad_residual or "",
        riesgo.impacto_residual or "",
        riesgo.severidad_residual if riesgo.residual_declarado else "",
        riesgo.disparador,
        panel.personas.get(riesgo.responsable_id or 0, ""),
        riesgo.revisar_el,
        panel.titulos.get(riesgo.mitigacion_task_id or 0, ""),
        ETIQUETA_ESTADO_RIESGO[riesgo.estado],
    ]
